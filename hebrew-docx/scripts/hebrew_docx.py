#!/usr/bin/env python3
"""hebrew_docx.py - Hebrew (right-to-left) Word documents that open correctly everywhere.

Usage:
    python3 hebrew_docx.py <input.md> <output.docx>      Markdown -> RTL .docx
    python3 hebrew_docx.py --fix <input.docx> <out.docx>  repair an existing Hebrew .docx
    python3 hebrew_docx.py --self-test

Requires python-docx (pip install python-docx).

Why this exists: Word lays every paragraph out left-to-right unless told otherwise, and it
takes the font, bold and size of Hebrew text from the complex-script properties
(w:rFonts/@w:cs, w:bCs, w:iCs, w:szCs), which python-docx never writes. A document built
without them opens with reversed tables, punctuation on the wrong side and Hebrew at the
wrong size. This script sets RTL on the document default, every paragraph, every run, every
section and every table, and mirrors bold/italic/size onto the complex-script twins.

Derived from the md2docx renderer of the legal-pleadings-il plugin (court rules removed).
"""
import html
import os
import re
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Arial"          # ships with macOS, Windows and Google Docs, and has Hebrew glyphs
CODE_FONT = "Courier New"
SIZE = 12
HEADING_PT = {1: 18, 2: 15, 3: 13, 4: 12}
LINE_SPACING = 1.25
SPACE_AFTER_PT = 6
MARGIN_CM = 2.5
INDENT_TWIPS = 425      # 0.75 cm per list level

# OOXML is order-sensitive: an element inserted after one of its schema successors is
# dropped silently by Word and Google Docs, and the paragraph renders LTR again. Every
# insertion below names the tags that must come after it.
_PPR_TAIL = ('w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
             'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
_PPR_AFTER_IND = ('w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap') + _PPR_TAIL
_PPR_AFTER_BIDI = ('w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind') + _PPR_AFTER_IND
_PPR_AFTER_JC = _PPR_TAIL[1:]
_PPR_AFTER_KEEPNEXT = ('w:keepLines', 'w:pageBreakBefore', 'w:framePr', 'w:widowControl',
                       'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd', 'w:tabs',
                       'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                       'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi') + _PPR_AFTER_BIDI
_RPR_AFTER_RTL = ('w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish', 'w:oMath',
                  'w:rPrChange')
_RPR_AFTER_LANG = _RPR_AFTER_RTL[3:]
_SECTPR_AFTER_BIDI = ('w:rtlGutter', 'w:docGrid', 'w:printerSettings', 'w:sectPrChange')
_TBLPR_AFTER_BORDERS = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
                        'w:tblDescription', 'w:tblPrChange')
_TBLPR_AFTER_BIDIVISUAL = ('w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
                           'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders') + _TBLPR_AFTER_BORDERS

# Invisible direction marks that chat-oriented Hebrew writing inserts. Inside a document
# they only fight the paragraph direction, so they never reach the .docx.
_BIDI_CONTROLS = re.compile('[‪-‮⁦-⁩]')
_INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
_LINK = re.compile(r'\[([^\]]+)\]\(([^)\s]+)\)')
_TABLE_SEP = re.compile(r'^\s*\|[\s:|-]+\|\s*$')
_HEADING = re.compile(r'^(#{1,6})\s+(.*)$')
_LIST = re.compile(r'^(\s*)([-*+]|\d{1,3}[.)])\s+(.*)$')
_HR = re.compile(r'^(-{3,}|\*{3,}|_{3,})$')


def _set_bidi(ppr, on=True):
    for old in ppr.findall(qn('w:bidi')):
        ppr.remove(old)
    el = OxmlElement('w:bidi')
    if not on:
        el.set(qn('w:val'), '0')
    ppr.insert_element_before(el, *_PPR_AFTER_BIDI)


def _set_jc(ppr, jc):
    for old in ppr.findall(qn('w:jc')):
        ppr.remove(old)
    if jc:
        el = OxmlElement('w:jc')
        el.set(qn('w:val'), jc)
        ppr.insert_element_before(el, *_PPR_AFTER_JC)


def _complex_script(rpr):
    """Mirror bold/italic/size onto the complex-script twins Word uses for Hebrew."""
    for latin, twin_tag in (('w:b', 'w:bCs'), ('w:i', 'w:iCs'), ('w:sz', 'w:szCs')):
        source = rpr.find(qn(latin))
        if source is None:
            continue
        for old in rpr.findall(qn(twin_tag)):
            rpr.remove(old)
        twin = OxmlElement(twin_tag)
        value = source.get(qn('w:val'))
        if value is not None:
            twin.set(qn('w:val'), value)
        source.addnext(twin)


def _rtl(par, jc=None):
    """Make one paragraph RTL. jc=None keeps the natural start (right) alignment.

    Only 'center' and 'both' are ever written: Google Docs reads 'left'/'right' logically
    in RTL while Word reads them physically, so either one lands on the wrong side
    somewhere."""
    ppr = par._p.get_or_add_pPr()
    _set_bidi(ppr)
    _set_jc(ppr, jc)
    for r in par._p.iter(qn('w:r')):   # includes runs inside hyperlinks
        rpr = r.get_or_add_rPr()
        _complex_script(rpr)
        for old in rpr.findall(qn('w:rtl')):
            rpr.remove(old)
        rpr.insert_element_before(OxmlElement('w:rtl'), *_RPR_AFTER_RTL)


def _rtl_sections(doc):
    for section in doc.sections:
        sect_pr = section._sectPr
        for old in sect_pr.findall(qn('w:bidi')):
            sect_pr.remove(old)
        sect_pr.insert_element_before(OxmlElement('w:bidi'), *_SECTPR_AFTER_BIDI)


def _rtl_table(table):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn('w:bidiVisual')):
        tbl_pr.remove(old)
    tbl_pr.insert_element_before(OxmlElement('w:bidiVisual'), *_TBLPR_AFTER_BIDIVISUAL)


def _style_rtl(doc):
    """Document-wide RTL defaults: Normal is RTL with a Hebrew-capable complex-script font,
    Hebrew proofing, and every style's bold/size mirrored for Hebrew."""
    normal = doc.styles['Normal']
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    if rfonts.get(qn('w:cs')) is None:
        rfonts.set(qn('w:cs'), FONT)
    rfonts.set(qn('w:hint'), 'cs')
    for old in rpr.findall(qn('w:lang')):
        rpr.remove(old)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:bidi'), 'he-IL')
    rpr.insert_element_before(lang, *_RPR_AFTER_LANG)
    _set_bidi(normal.element.get_or_add_pPr())
    for style in doc.styles:
        style_rpr = getattr(style.element, 'rPr', None)
        if style_rpr is not None:
            _complex_script(style_rpr)


def _style_look(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(SIZE)
    rpr = normal.element.get_or_add_rPr()
    rpr.find(qn('w:rFonts')).set(qn('w:cs'), FONT)
    _complex_script(rpr)
    fmt = normal.paragraph_format
    fmt.line_spacing = LINE_SPACING
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(SPACE_AFTER_PT)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(MARGIN_CM)
        section.left_margin = section.right_margin = Cm(MARGIN_CM)


def _indent(par, start, hanging=0):
    """w:left is the START side (right) in an RTL paragraph. Never add w:start as well:
    Google Docs sums the two on import."""
    ppr = par._p.get_or_add_pPr()
    for old in ppr.findall(qn('w:ind')):
        ppr.remove(old)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str(start))
    if hanging:
        ind.set(qn('w:hanging'), str(hanging))
    ppr.insert_element_before(ind, *_PPR_AFTER_IND)


def _keep_with_next(par):
    """A heading never ends a page while its section starts on the next one."""
    ppr = par._p.get_or_add_pPr()
    if not ppr.findall(qn('w:keepNext')):
        ppr.insert_element_before(OxmlElement('w:keepNext'), *_PPR_AFTER_KEEPNEXT)


# In an RTL paragraph the hyphen of "5%-15%" resolves as a neutral (UAX #9 applies W4
# before W5, so the % blocks the number-separator rule) and the range displays as
# "15%-5%" - in Word, Pages and browsers alike. LEFT-TO-RIGHT MARKs pin it.
_RANGE_HYPHEN = re.compile(r'(?<=[\d%])([-–])(?=\d)')


def _clean(text):
    text = html.unescape(text)
    text = re.sub(r'\\([\\`*_{}\[\]()#+\-.!>~|])', r'\1', text)
    return _RANGE_HYPHEN.sub('‎\\1‎', text)


def _add_runs(par, text):
    text = _LINK.sub(lambda m: m.group(1) if m.group(1) == m.group(2)
                     else f"{m.group(1)} ({m.group(2)})", text)
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            par.add_run(_clean(part[2:-2])).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            par.add_run(_clean(part[1:-1])).italic = True
        elif part.startswith('`') and part.endswith('`'):
            par.add_run(_clean(part[1:-1]))
        else:
            par.add_run(_clean(part))   # python-docx turns \t into a real <w:tab/>


def _table_borders(table):
    """Borders on the table itself: Google Docs drops table styles it did not write."""
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        line = OxmlElement(f'w:{edge}')
        for key, value in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '0'),
                           ('w:color', '808080')):
            line.set(qn(key), value)
        borders.append(line)
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn('w:tblBorders')):
        tbl_pr.remove(old)
    tbl_pr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)


def _add_table(doc, block):
    rows = [[c.strip() for c in ln.strip().strip('|').split('|')]
            for ln in block if not _TABLE_SEP.match(ln)]
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    _rtl_table(table)          # first column on the right
    _table_borders(table)
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            par = cell.paragraphs[0]
            _add_runs(par, row[ci] if ci < len(row) else '')
            par.paragraph_format.space_after = Pt(2)
            if ri == 0:
                for run in par.runs:
                    run.bold = True
                shd = OxmlElement('w:shd')
                for key, value in (('w:val', 'clear'), ('w:color', 'auto'), ('w:fill', 'F2F2F2')):
                    shd.set(qn(key), value)
                cell._tc.get_or_add_tcPr().append(shd)
            _rtl(par)


def _code_paragraph(doc, line):
    par = doc.add_paragraph()
    run = par.add_run(line if line else ' ')
    run.font.name = CODE_FONT
    run.font.size = Pt(10)
    par.paragraph_format.space_after = Pt(0)
    _set_bidi(par._p.get_or_add_pPr(), on=False)   # code reads left-to-right


def _footer_page_numbers(doc):
    for section in doc.sections:
        par = section.footer.paragraphs[0]
        for tag, text in (('begin', None), ('instr', ' PAGE '), ('separate', None),
                          ('text', '1'), ('end', None)):
            run = par.add_run()
            run.font.size = Pt(10)
            if tag == 'instr':
                el = OxmlElement('w:instrText')
                el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                el.text = text
            elif tag == 'text':
                el = OxmlElement('w:t')
                el.text = text
            else:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), tag)
            run._r.append(el)
        _rtl(par, jc='center')


def render(md_text, out_path):
    doc = Document()
    _style_rtl(doc)
    _style_look(doc)
    lines = _BIDI_CONTROLS.sub('', md_text).splitlines()
    i, in_code = 0, False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        i += 1
        if stripped.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            _code_paragraph(doc, line)
            continue
        if not stripped or _HR.match(stripped) or stripped.startswith('<!--'):
            continue
        if stripped.startswith('|') and i < len(lines) and _TABLE_SEP.match(lines[i]):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        par = doc.add_paragraph()
        heading = _HEADING.match(stripped)
        listed = _LIST.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            _add_runs(par, heading.group(2))
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(HEADING_PT[level])
            par.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            _rtl(par, jc='center' if level == 1 else None)
            _keep_with_next(par)
        elif stripped.startswith('>'):
            _add_runs(par, stripped.lstrip('>').strip())
            for run in par.runs:
                run.italic = True
            _indent(par, INDENT_TWIPS)
            _rtl(par)
        elif listed:
            indent, marker, text = listed.groups()
            level = len(indent.expandtabs(4)) // 2
            _add_runs(par, ('•' if marker in '-*+' else marker) + '\t' + text)
            _indent(par, INDENT_TWIPS * (level + 1), INDENT_TWIPS)
            par.paragraph_format.space_after = Pt(2)
            _rtl(par)
        else:
            _add_runs(par, stripped)
            _rtl(par)
    body = list(doc.element.body)
    if len(body) > 1 and body[-2].tag == qn('w:tbl'):
        _rtl(doc.add_paragraph())        # a table may not be the last body element
    _rtl_sections(doc)
    _footer_page_numbers(doc)
    doc.save(out_path)
    return out_path


def _walk(container, on_par, on_table):
    for par in container.paragraphs:
        on_par(par)
    for table in container.tables:
        on_table(table)
        for row in table.rows:
            for cell in row.cells:       # merged cells repeat; every step is idempotent
                _walk(cell, on_par, on_table)


def fix(in_path, out_path):
    """Make an existing Hebrew .docx RTL without changing its content or look."""
    doc = Document(in_path)
    _style_rtl(doc)

    def fix_par(par):
        jc_el = par._p.pPr.find(qn('w:jc')) if par._p.pPr is not None else None
        keep = jc_el.get(qn('w:val')) if jc_el is not None else None
        _rtl(par, jc=keep if keep in ('center', 'both', 'distribute') else None)
        for t in par._p.iter(qn('w:t')):   # invisible marks only; the text itself is kept
            if t.text:
                t.text = _RANGE_HYPHEN.sub('‎\\1‎', t.text)

    _walk(doc, fix_par, _rtl_table)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:   # touching a linked part would create one
                _walk(part, fix_par, _rtl_table)
    _rtl_sections(doc)
    doc.save(out_path)
    return out_path


def self_test():
    sample = "\n".join([
        "# מסמך בדיקה",
        "פסקה עם **הדגשה** ומילה באנגלית MVP, וקישור [אתר](https://example.com).",
        "⁦תו בידוד⁩ שלא אמור להופיע",
        "## רשימה",
        "- פריט ראשון",
        "  - פריט מקונן",
        "1. צעד ראשון",
        "> ציטוט",
        "טווח 5%-15% בחודש",
        "---",
        "```",
        "print('hello')",
        "```",
        "| עמודה א | עמודה ב |",
        "|---|---|",
        "| 1 | שתיים |",
    ])
    out = os.path.join(tempfile.gettempdir(), 'hebrew_docx_selftest.docx')
    render(sample, out)
    d = Document(out)
    normal = d.styles['Normal'].element.xml
    assert '<w:bidi/>' in normal, "Normal style must be RTL"
    assert 'w:cs="%s"' % FONT in normal and 'w:szCs' in normal, "Normal needs a Hebrew font and size"
    assert 'w:bidi="he-IL"' in normal, "Normal needs Hebrew proofing language"
    assert '<w:bidi/>' in d.sections[0]._sectPr.xml, "section must be RTL"
    code = [p for p in d.paragraphs if p.text == "print('hello')"]
    assert code and 'w:bidi w:val="0"' in code[0]._p.xml, "code must stay left-to-right"
    for par in d.paragraphs:
        if par.text == "print('hello')":   # python-docx builds a new Paragraph per access
            continue
        x = par._p.xml
        assert '<w:bidi/>' in x, f"paragraph not RTL: {par.text!r}"
        assert 'w:val="left"' not in x and 'w:val="right"' not in x, "jc left/right renders on the wrong side somewhere"
        for r in par._p.iter(qn('w:r')):
            rx = r.xml
            assert '<w:rtl/>' in rx, "every run must be RTL"
            if '<w:b/>' in rx:
                assert '<w:bCs/>' in rx, "Hebrew bold needs w:bCs"
            if '<w:sz ' in rx:
                assert '<w:szCs ' in rx, "Hebrew size needs w:szCs"
    heading = d.paragraphs[0]
    assert 'w:val="center"' in heading._p.xml and 'w:keepNext' in heading._p.xml
    text = "\n".join(p.text for p in d.paragraphs)
    assert not _BIDI_CONTROLS.search(text), "bidi control characters leaked into the document"
    assert "(https://example.com)" in text, "link target lost"
    assert "5%‎-‎15%" in text, "a number range must be pinned with LRM marks"
    bullet = next(p for p in d.paragraphs if p.text.startswith('•'))
    assert '<w:tab/>' in bullet._p.xml and 'w:hanging="%d"' % INDENT_TWIPS in bullet._p.xml
    nested = [p for p in d.paragraphs if p.text.startswith('•')][1]
    assert 'w:left="%d"' % (2 * INDENT_TWIPS) in nested._p.xml, "nested item must indent one more level"
    assert '---' not in text, "a horizontal rule must not render as text"
    assert len(d.tables) == 1
    tbl = d.tables[0]._tbl.tblPr.xml
    assert 'w:bidiVisual' in tbl and 'w:tblBorders' in tbl, "table must be RTL with its own borders"
    assert d.element.body[-2].tag == qn('w:p'), "a table may not be the last body element"
    footer = d.sections[0].footer.paragraphs[0]._p.xml
    assert 'PAGE' in footer and 'w:fldCharType="end"' in footer, "footer page number missing"

    # --fix on a plain LTR document built the way generic tools build it.
    src = os.path.join(tempfile.gettempdir(), 'hebrew_docx_ltr.docx')
    ltr = Document()
    p = ltr.add_paragraph()
    p.add_run("שלום").bold = True
    p.alignment = 0      # explicit LEFT
    c = ltr.add_paragraph("ממורכז")
    c.alignment = 1      # CENTER must survive
    ltr.add_paragraph("כ-4%–5%")
    t = ltr.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "א"
    ltr.save(src)
    fixed = Document(fix(src, os.path.join(tempfile.gettempdir(), 'hebrew_docx_fixed.docx')))
    x0, x1 = fixed.paragraphs[0]._p.xml, fixed.paragraphs[1]._p.xml
    assert '<w:bidi/>' in x0 and '<w:rtl/>' in x0 and '<w:bCs/>' in x0
    assert 'w:val="left"' not in x0, "a forced LEFT alignment must go"
    assert 'w:val="center"' in x1, "centered text must stay centered"
    assert "4%‎–‎5%" in fixed.paragraphs[2].text, "--fix must pin en-dash ranges too"
    assert 'w:bidiVisual' in fixed.tables[0]._tbl.tblPr.xml
    assert '<w:bidi/>' in fixed.tables[0].cell(0, 0).paragraphs[0]._p.xml
    assert '<w:bidi/>' in fixed.sections[0]._sectPr.xml
    print(f"hebrew_docx self-test OK ({out})")


if __name__ == '__main__':
    args = sys.argv[1:]
    if args == ['--self-test']:
        self_test()
    elif len(args) == 3 and args[0] == '--fix':
        print(f"wrote {fix(args[1], args[2])}")
    elif len(args) == 2:
        with open(args[0], encoding='utf-8') as f:
            print(f"wrote {render(f.read(), args[1])}")
    else:
        print(__doc__)
        sys.exit(1)
